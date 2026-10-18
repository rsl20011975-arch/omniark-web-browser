"""Read-only document extraction; never executes document macros or formulas."""
import json
import sys
import zipfile
from pathlib import Path

MAX_TEXT = 60_000
MAX_ARCHIVE_BYTES = 50_000_000

def check_archive(source):
    with zipfile.ZipFile(source) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_ARCHIVE_BYTES:
            raise ValueError('Document archive exceeds extraction limit')

def extract(source):
    suffix = source.suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        document = PdfReader(source)
        if document.is_encrypted:
            raise ValueError('Encrypted PDF requires manual decryption')
        parts = [page.extract_text() or '' for page in document.pages]
        return '\n\n'.join(parts), len(parts), ['text extraction; page layout not preserved']
    check_archive(source)
    if suffix == '.docx':
        from docx import Document
        document = Document(source)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.extend('\t'.join(cell.text for cell in row.cells) for row in table.rows)
        return '\n'.join(parts), len(document.paragraphs), ['embedded images are not OCR processed']
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(source, read_only=True, data_only=True)
        parts = []
        try:
            for sheet in book:
                parts.append('## ' + sheet.title)
                parts.extend('\t'.join('' if value is None else str(value) for value in row)
                             for row in sheet.iter_rows(values_only=True))
            return '\n'.join(parts), len(book.sheetnames), ['cached formula values only; no recalculation']
        finally:
            book.close()
    raise ValueError('Unsupported document type')

try:
    source = Path(sys.argv[1])
    text, units, warnings = extract(source)
    if not text.strip():
        raise ValueError('No readable text. Scanned documents require OCR.')
    print(json.dumps({'text': text[:MAX_TEXT], 'units': units, 'warnings': warnings,
                      'truncated': len(text) > MAX_TEXT}, ensure_ascii=False))
except Exception as error:
    print(str(error), file=sys.stderr)
    sys.exit(1)
