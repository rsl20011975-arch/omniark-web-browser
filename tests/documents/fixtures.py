"""Small extraction regression fixtures, not user-facing document deliverables."""
from pathlib import Path
from urllib.request import urlopen
from docx import Document
from openpyxl import Workbook

root = Path(__file__).resolve().parents[2] / 'data' / 'document-fixtures'
root.mkdir(parents=True, exist_ok=True)
document = Document()
document.add_paragraph('Workbench document extraction fixture')
table = document.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = 'Project'
table.rows[0].cells[1].text = 'Quantity'
for index in range(1, 11):
    row = table.add_row()
    row.cells[0].text = f'Project {index}'
    row.cells[1].text = str(index * 3)
document.save(root / 'projects.docx')
book = Workbook()
sheet = book.active
sheet.title = 'Projects'
sheet.append(['Project', 'Quantity'])
for index in range(1, 11):
    sheet.append([f'Project {index}', index * 3])
book.save(root / 'projects.xlsx')
with urlopen('https://www.w3.org/WAI/ER/tests/xhtml/testfiles/resources/pdf/dummy.pdf', timeout=30) as response:
    (root / 'dummy.pdf').write_bytes(response.read())
print(root)
